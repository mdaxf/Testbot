"""Builds test_cases/ad_demo/ad_demo_approval_workflow.xlsx -- a 3-case,
3-user, end-to-end approval-workflow suite for the AD_demo tenant, following
the exact workbook layout documented in README.md / examples/excel_demo.xlsx
(Suite, Index, CommonData, one sheet per CaseID).

Run once to (re)generate the workbook:
    python scripts/build_ad_demo_suite.py

Design notes (read before editing the step tables below):

- No `label` or `ai` TargetStrategy anywhere in this suite. The Portal
  login form's <label>/<input> pairs are plain siblings with no for/id/
  aria association, so Playwright's native get_by_label always resolves 0
  matches there; the framework's "ai" fallback then requires a configured
  AI_VISION_PROVIDER key, which this testing project's .env does not set.
  Login fields are targeted by `placeholder` (Tenant has placeholder=
  "default") or `css` on the `autocomplete` attribute (Username/Password
  have autocomplete="username"/"current-password" but no placeholder) --
  both resolve natively, no AI call needed anywhere in this suite.
- The AI-recommended action card in TC-001 is targeted by `xpath`, scoped
  off the card's own machine-rendered `-> {target_id}` line (e.g.
  "plm-bom-change-impact") rather than the LLM-paraphrased recommendation
  title/body -- the skill can only recommend a flow/action id from its
  configured allowlist, so that id string is stable across runs even
  though the surrounding wording is not.
- TC-002/TC-003 assume the request created by TC-001 is the ONLY item in
  the approver's "assigned" inbox at that moment (a clean-state
  assumption) and target the first table row generically via CSS.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.worksheet import Worksheet

OUT_PATH = Path(__file__).resolve().parent.parent / "test_cases" / "ad_demo" / "ad_demo_approval_workflow.xlsx"

HEADER_FONT = Font(name="Calibri", bold=True)
STEP_COLUMNS = [
    "StepNo", "Description", "Action", "TargetStrategy", "TargetValue", "TargetName",
    "Input", "Query", "Connection", "ExpectedType", "ExpectedValue",
    "ExpectedTargetStrategy", "ExpectedTargetValue", "ExpectedColumn",
    "CaptureVar", "CaptureFrom", "CaptureTargetStrategy", "CaptureTargetValue",
    "CaptureColumn", "DelayAfterMs",
]


def _write_header_row(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = HEADER_FONT


def _write_rows(ws: Worksheet, start_row: int, headers: list[str], rows: list[dict]) -> None:
    for r, row_data in enumerate(rows, start=start_row):
        for col, header in enumerate(headers, start=1):
            value = row_data.get(header)
            if value is not None:
                ws.cell(row=r, column=col, value=value)


def build_suite_sheet(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("Suite")
    headers = ["SuiteID", "SuiteName", "Environment", "OnCaseFail", "DefaultStepDelayMs"]
    _write_header_row(ws, 1, headers)
    _write_rows(ws, 2, headers, [{
        "SuiteID": "AD-DEMO-APPROVAL",
        "SuiteName": "AD_demo: launch dashboard, approve AI action, multi-user multi-stage approval",
        "Environment": "ad_demo_portal",
        "OnCaseFail": "continue",
        "DefaultStepDelayMs": 1500,
    }])
    ws.column_dimensions["B"].width = 55


def build_index_sheet(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("Index")
    headers = ["Sequence", "CaseID", "Run"]
    _write_header_row(ws, 1, headers)
    _write_rows(ws, 2, headers, [
        {"Sequence": 1, "CaseID": "TC-001", "Run": "Y"},
        {"Sequence": 2, "CaseID": "TC-002", "Run": "Y"},
        {"Sequence": 3, "CaseID": "TC-003", "Run": "Y"},
    ])


def build_common_data_sheet(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("CommonData")
    headers = ["VarName", "Source", "Value"]
    _write_header_row(ws, 1, headers)
    # Three distinct AD_demo logins, one per approval-chain role, so the
    # suite genuinely exercises "different users at different stages"
    # rather than one superuser account clicking through everything.
    # Created via ConfigStore.create_user -- see project chat history for
    # the exact provisioning script; re-run it if these accounts don't
    # exist yet in the target's mcp.db.
    _write_rows(ws, 2, headers, [
        {"VarName": "Tenant", "Source": "constant", "Value": "AD_demo"},
        {"VarName": "RequesterUsername", "Source": "constant", "Value": "ad_requester"},
        {"VarName": "RequesterPassword", "Source": "constant", "Value": "TestPass123!"},
        {"VarName": "EngineerUsername", "Source": "constant", "Value": "ad_engineer"},
        {"VarName": "EngineerPassword", "Source": "constant", "Value": "TestPass123!"},
        {"VarName": "PmUsername", "Source": "constant", "Value": "ad_pm"},
        {"VarName": "PmPassword", "Source": "constant", "Value": "TestPass123!"},
    ])
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["C"].width = 20


def _login_steps(start_no: int, username_var: str, password_var: str) -> list[dict]:
    """Common login sequence -- Tenant/Username/Password + Sign in, asserting
    the post-login Studio shell rendered (its "DASHBOARDS" nav section header
    is always present regardless of which page HomeRedirect lands on)."""
    return [
        {
            "StepNo": start_no, "Description": "Open the Portal login page",
            "Action": "navigate", "TargetStrategy": "url", "TargetValue": "{base_url}/",
        },
        {
            "StepNo": start_no + 1, "Description": "Enter tenant id",
            "Action": "type", "TargetStrategy": "placeholder", "TargetValue": "default",
            "Input": "{Tenant}",
        },
        {
            "StepNo": start_no + 2, "Description": "Enter username",
            "Action": "type", "TargetStrategy": "css", "TargetValue": 'input[autocomplete="username"]',
            "Input": f"{{{username_var}}}",
        },
        {
            "StepNo": start_no + 3, "Description": "Enter password",
            "Action": "type", "TargetStrategy": "css", "TargetValue": 'input[autocomplete="current-password"]',
            "Input": f"{{{password_var}}}",
        },
        {
            "StepNo": start_no + 4,
            "Description": (
                "Click the Sign in submit button (css, not role/name='Sign in': that also "
                "matches the 'Sign in' TAB toggle above the form, which is ambiguous)"
            ),
            "Action": "click", "TargetStrategy": "css", "TargetValue": "button.btn-primary",
            "ExpectedType": "visible", "ExpectedTargetStrategy": "text", "ExpectedTargetValue": "DASHBOARDS",
            "DelayAfterMs": 20000,
        },
    ]


def build_tc001(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("TC-001")
    ws["A1"] = "Title"
    ws["B1"] = "Requester launches the Engineering dashboard and approves the AI-recommended PLM/BOM Change Impact action"
    ws["A2"] = "AreaPath"
    ws["B2"] = "AD_demo/ApprovalWorkflow"
    ws["A3"] = "Priority"
    ws["B3"] = 1
    ws["A4"] = "Preconditions"
    ws["B4"] = (
        "mom-mcp (8787, serving its own embedded Studio at /admin/studio -- or "
        "whatever --env ad_demo_portal points at) is running; AD_demo's and_engineering "
        "dashboard action_panel widget "
        "recommends plm-bom-change-impact (gated by the engineering_change_approval "
        "policy: stage 1 = engineering, stage 2 = program_manager); ad_requester holds "
        "only the 'executive' role (can view the dashboard, cannot approve either stage)."
    )

    header_row = 6
    ws.cell(row=header_row, column=1, value="StepNo")
    _write_header_row(ws, header_row, STEP_COLUMNS)

    steps = _login_steps(1, "RequesterUsername", "RequesterPassword")
    steps += [
        {
            "StepNo": 6, "Description": "Open the Engineering & PLM Change Management Cockpit dashboard",
            "Action": "navigate", "TargetStrategy": "url",
            "TargetValue": "{base_url}/studio/dashboard/and_engineering",
        },
        {
            "StepNo": 7,
            "Description": (
                "Wait for the whole dashboard to finish loading -- 12 widgets including "
                "two separate LLM-backed calls (ai_insight_panel, action_panel), confirmed "
                "by a live run to still show 'Loading...' on the basic stat cards at 12s"
            ),
            "Action": "wait", "Input": "45000",
        },
        {
            "StepNo": 8,
            "Description": (
                "Click Approve on the recommendation card whose target is plm-bom-change-impact "
                "(scoped via the card's own machine-rendered target-id line, not the LLM-paraphrased "
                "title/body, since only the id is guaranteed stable across runs)"
            ),
            "Action": "click", "TargetStrategy": "xpath",
            "TargetValue": '//p[contains(., "plm-bom-change-impact")]/parent::div//button[normalize-space()="Approve"]',
            "ExpectedType": "visible", "ExpectedTargetStrategy": "role",
            "ExpectedTargetValue": "button",
            "ExpectedColumn": None,
        },
        {
            "StepNo": 9, "Description": "Confirm execution in the Yes, execute / Cancel prompt",
            "Action": "click", "TargetStrategy": "role", "TargetValue": "button", "TargetName": "Yes, execute",
            "DelayAfterMs": 3000,
            "ExpectedType": "text_contains", "ExpectedValue": "Submitted for approval",
            "ExpectedTargetStrategy": "text", "ExpectedTargetValue": "Submitted for approval",
        },
    ]
    _write_rows(ws, header_row + 1, STEP_COLUMNS, steps)
    for col_letter, width in {"B": 60, "E": 55}.items():
        ws.column_dimensions[col_letter].width = width


def build_tc002(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("TC-002")
    ws["A1"] = "Title"
    ws["B1"] = "Engineering approver approves stage 1 of 2, advancing the request to the Program Manager stage"
    ws["A2"] = "AreaPath"
    ws["B2"] = "AD_demo/ApprovalWorkflow"
    ws["A3"] = "Priority"
    ws["B3"] = 1
    ws["A4"] = "Preconditions"
    ws["B4"] = (
        "TC-001 has already run and created a pending engineering_change_approval "
        "request; ad_engineer holds the 'engineering' role and this is the ONLY item "
        "in their assigned inbox (clean-state assumption -- see suite notes)."
    )

    header_row = 6
    _write_header_row(ws, header_row, STEP_COLUMNS)

    steps = _login_steps(1, "EngineerUsername", "EngineerPassword")
    steps += [
        {
            "StepNo": 6, "Description": "Open the assigned-to-me Inbox tab",
            "Action": "navigate", "TargetStrategy": "url", "TargetValue": "{base_url}/studio/inbox/assigned",
        },
        {
            "StepNo": 7, "Description": "Wait for the assigned-requests table to load",
            "Action": "wait", "Input": "25000",
        },
        {
            "StepNo": 8, "Description": "Open the pending request (first, and only, row)",
            "Action": "click", "TargetStrategy": "css", "TargetValue": "table tbody tr:first-child",
            "DelayAfterMs": 1500,
            "ExpectedType": "visible", "ExpectedTargetStrategy": "text", "ExpectedTargetValue": "Approval chain",
        },
        {
            "StepNo": 9, "Description": "Select the pending engineering stage node in the approval-chain canvas",
            "Action": "click", "TargetStrategy": "text", "TargetValue": "engineering",
            "ExpectedType": "visible", "ExpectedTargetStrategy": "role", "ExpectedTargetValue": "button",
        },
        {
            "StepNo": 10, "Description": "Click Approve for this stage",
            "Action": "click", "TargetStrategy": "role", "TargetValue": "button", "TargetName": "Approve",
            "DelayAfterMs": 2000,
            "ExpectedType": "visible", "ExpectedTargetStrategy": "text", "ExpectedTargetValue": "approved",
        },
    ]
    _write_rows(ws, header_row + 1, STEP_COLUMNS, steps)
    for col_letter, width in {"B": 60, "E": 55}.items():
        ws.column_dimensions[col_letter].width = width


def build_tc003(wb: openpyxl.Workbook) -> None:
    ws = wb.create_sheet("TC-003")
    ws["A1"] = "Title"
    ws["B1"] = "Program Manager approves the final stage, and the flow actually executes"
    ws["A2"] = "AreaPath"
    ws["B2"] = "AD_demo/ApprovalWorkflow"
    ws["A3"] = "Priority"
    ws["B3"] = 1
    ws["A4"] = "Preconditions"
    ws["B4"] = (
        "TC-002 has already run and advanced the request to stage 2 (program_manager); "
        "ad_pm holds the 'program_manager' role and this is the ONLY item in their "
        "assigned inbox (clean-state assumption -- see suite notes)."
    )

    header_row = 6
    _write_header_row(ws, header_row, STEP_COLUMNS)

    steps = _login_steps(1, "PmUsername", "PmPassword")
    steps += [
        {
            "StepNo": 6, "Description": "Open the assigned-to-me Inbox tab",
            "Action": "navigate", "TargetStrategy": "url", "TargetValue": "{base_url}/studio/inbox/assigned",
        },
        {
            "StepNo": 7, "Description": "Wait for the assigned-requests table to load",
            "Action": "wait", "Input": "25000",
        },
        {
            "StepNo": 8, "Description": "Open the pending request (first, and only, row)",
            "Action": "click", "TargetStrategy": "css", "TargetValue": "table tbody tr:first-child",
            "DelayAfterMs": 1500,
            "ExpectedType": "visible", "ExpectedTargetStrategy": "text", "ExpectedTargetValue": "Approval chain",
        },
        {
            "StepNo": 9, "Description": "Select the pending program_manager stage node in the approval-chain canvas",
            "Action": "click", "TargetStrategy": "text", "TargetValue": "program_manager",
            "ExpectedType": "visible", "ExpectedTargetStrategy": "role", "ExpectedTargetValue": "button",
        },
        {
            "StepNo": 10, "Description": "Click Approve for the final stage -- this is the last approval, so the flow executes immediately",
            "Action": "click", "TargetStrategy": "role", "TargetValue": "button", "TargetName": "Approve",
            "DelayAfterMs": 4000,
            "ExpectedType": "visible", "ExpectedTargetStrategy": "text", "ExpectedTargetValue": "Execution result",
        },
    ]
    _write_rows(ws, header_row + 1, STEP_COLUMNS, steps)
    for col_letter, width in {"B": 70, "E": 55}.items():
        ws.column_dimensions[col_letter].width = width


def main() -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # drop the default empty sheet
    build_suite_sheet(wb)
    build_index_sheet(wb)
    build_common_data_sheet(wb)
    build_tc001(wb)
    build_tc002(wb)
    build_tc003(wb)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
