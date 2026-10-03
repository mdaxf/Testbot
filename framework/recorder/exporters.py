from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from framework.models import TestSuite


def _clean_steps(steps: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out = []
    for i, step in enumerate(steps, start=1):
        clean = {k: v for k, v in step.items() if not k.startswith("_")}
        out.append({"step_no": i, **clean})
    return out


def build_suite(session: Any, *, suite_id: str, suite_name: str, case_id: str, title: str) -> TestSuite:
    """Snapshot of a RecordingSession as a validated TestSuite (so anything exported is
    guaranteed loadable by the runner)."""
    with session.lock:
        steps = _clean_steps(session.steps)
        variables = {k: dict(v) for k, v in session.variables.items()}
        base_url = session.base_url
    data: dict[str, Any] = {
        "suite_id": suite_id, "suite_name": suite_name, "on_case_fail": "continue",
        "cases": [{"id": case_id, "title": title, "steps": steps}],
    }
    if base_url:
        data["base_url"] = base_url
    if variables:
        data["variables"] = variables
    return TestSuite.model_validate(data)


def _suite_json_dict(suite: TestSuite, session: Any) -> dict[str, Any]:
    """TestSuite merges shared variables into each case on validation; for a readable file we
    write them once at suite level again."""
    data = suite.model_dump(mode="json", exclude_none=True)
    data.pop("variables", None)
    data.pop("connections", None)
    with session.lock:
        variables = {k: dict(v) for k, v in session.variables.items()}
    for case in data["cases"]:
        case.pop("variables", None)
        for step in case["steps"]:
            for key in [k for k, v in step.items() if v in (None, {}, "") and k not in ("input",)]:
                step.pop(key)
            step.pop("connection", None)
            if step.get("timeout_ms") == 10_000:
                step.pop("timeout_ms")
    ordered: dict[str, Any] = {"suite_id": data["suite_id"], "suite_name": data["suite_name"]}
    if data.get("base_url"):
        ordered["base_url"] = data["base_url"]
    if variables:
        ordered["variables"] = variables
    ordered["on_case_fail"] = data["on_case_fail"]
    ordered["cases"] = data["cases"]
    return ordered


def write_json(suite: TestSuite, session: Any, path: Path) -> Path:
    path.write_text(json.dumps(_suite_json_dict(suite, session), indent=2, ensure_ascii=False), encoding="utf-8")
    return path


def write_excel(suite: TestSuite, session: Any, path: Path) -> Path:
    from openpyxl import Workbook

    case = suite.cases[0]
    wb = Workbook()
    ws = wb.active
    ws.title = "Suite"
    ws.append(["SuiteID", "SuiteName", "BaseUrl", "OnCaseFail", "DefaultStepDelayMs"])
    ws.append([suite.suite_id, suite.suite_name, suite.base_url or "", suite.on_case_fail, 0])

    idx = wb.create_sheet("Index")
    idx.append(["Sequence", "CaseID", "Run"])
    idx.append([1, case.id, "Y"])

    with session.lock:
        variables = {k: dict(v) for k, v in session.variables.items()}
    if variables:
        common = wb.create_sheet("CommonData")
        common.append(["VarName", "Source", "Value"])
        for name, var in variables.items():
            common.append([name, var["source"], var.get("value")])

    sheet = wb.create_sheet(case.id[:31])
    sheet.append(["Title", case.title])
    sheet.append([])
    sheet.append(["StepNo", "Description", "Action", "TargetStrategy", "TargetValue", "TargetName", "Input",
                  "ExpectedType", "ExpectedValue", "TimeoutMs", "Config"])
    for step in case.steps:
        target, expected = step.target, step.expected
        sheet.append([
            step.step_no, step.description, step.action,
            target.strategy if target else None, target.value if target else None, target.name if target else None,
            step.input, expected.type if expected else None, expected.value if expected else None,
            step.timeout_ms if step.timeout_ms != 10_000 else None,
            json.dumps(step.config) if step.config else None,
        ])
    wb.save(path)
    return path


# ------------------------------------------------------------------ Playwright script

def _locator_code(target: Any) -> str:
    s, v, n = target.strategy, target.value, target.name
    if s == "role":
        code = f"page.get_by_role({v!r}, name={n!r})" if n else f"page.get_by_role({v!r})"
    elif s == "label":
        code = f"page.get_by_label({v!r})"
    elif s == "placeholder":
        code = f"page.get_by_placeholder({v!r})"
    elif s == "text":
        code = f"page.get_by_text({v!r})"
    elif s == "testid":
        code = f"page.get_by_test_id({v!r})"
    elif s == "xpath":
        code = f"page.locator({('xpath=' + v)!r})"
    else:
        code = f"page.locator({v!r})"
    return f"{code}.nth({target.nth - 1})" if target.nth else code


def _value_code(text: str) -> str:
    """A step input/url may contain {var} placeholders -> resolve them from the VARS dict."""
    if re.search(r"\{[A-Za-z_]\w*\}", text):
        return f"_fmt({text!r})"
    return repr(text)


def write_playwright_script(suite: TestSuite, session: Any, path: Path) -> Path:
    case = suite.cases[0]
    with session.lock:
        variables = {k: dict(v) for k, v in session.variables.items()}
    lines = [
        f'"""Recorded by testbot: {suite.suite_name} -- {case.title}',
        'Run with: pip install playwright && playwright install chromium && python <this file>',
        '"""',
        "import os",
        "import re",
        "",
        "from playwright.sync_api import expect, sync_playwright",
        "",
        f"BASE_URL = {suite.base_url or ''!r}",
        "VARS = {",
    ]
    for name, var in variables.items():
        lines.append(f"    {name!r}: os.environ.get({name.upper()!r}, {var.get('value')!r}),")
    lines += [
        "}",
        "",
        "",
        "def _fmt(text):",
        '    return text.format(base_url=BASE_URL, **VARS)',
        "",
        "",
        "def run(page):",
    ]
    for step in case.steps:
        lines.append(f"    # {step.step_no}. {step.description}")
        a, t = step.action, step.target
        if a == "navigate":
            lines.append(f"    page.goto({_value_code(t.value)})")
        elif a == "click":
            lines.append(f"    {_locator_code(t)}.click()")
        elif a == "type":
            lines.append(f"    {_locator_code(t)}.fill({_value_code(step.input or '')})")
        elif a == "select":
            lines.append(f"    {_locator_code(t)}.select_option({_value_code(step.input or '')})")
        elif a == "press":
            lines.append(f"    {_locator_code(t)}.press({step.input!r})")
        elif a == "upload":
            lines.append(f"    {_locator_code(t)}.set_input_files({step.input!r})  # TODO: set a real file path")
        elif a == "wait_until":
            cond, _, arg = (step.input or "hidden").partition(":")
            loc, ms = _locator_code(t) + ".first", step.timeout_ms
            call = {
                "hidden": f"to_be_hidden(timeout={ms})", "visible": f"to_be_visible(timeout={ms})",
                "value": f"to_have_value({arg!r}, timeout={ms})", "has_value": f"to_have_value(re.compile('.+'), timeout={ms})",
                "text": f"to_have_text({arg!r}, timeout={ms})", "text_contains": f"to_contain_text({arg!r}, timeout={ms})",
            }[cond]
            lines.append(f"    expect({loc}).{call}")
        else:
            lines.append(f"    # (unsupported action in script export: {a})")
        if step.expected and step.expected.type == "url_contains":
            lines.append(f"    expect(page).to_have_url(re.compile(re.escape({step.expected.value!r})))")
    lines += [
        "",
        "",
        'if __name__ == "__main__":',
        "    with sync_playwright() as pw:",
        "        browser = pw.chromium.launch(headless=False)",
        "        page = browser.new_context().new_page()",
        "        try:",
        "            run(page)",
        "        finally:",
        "            browser.close()",
        "",
    ]
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def export_all(session: Any, out_dir: Path, *, suite_id: str, suite_name: str, case_id: str, title: str,
               formats: set[str]) -> list[Path]:
    """formats: any of {'json', 'excel', 'playwright'}. Returns the files written."""
    out_dir.mkdir(parents=True, exist_ok=True)
    suite = build_suite(session, suite_id=suite_id, suite_name=suite_name, case_id=case_id, title=title)
    written: list[Path] = []
    if "json" in formats:
        written.append(write_json(suite, session, out_dir / f"{suite_id}.json"))
    if "excel" in formats:
        written.append(write_excel(suite, session, out_dir / f"{suite_id}.xlsx"))
    if "playwright" in formats:
        written.append(write_playwright_script(suite, session, out_dir / f"{suite_id}_playwright.py"))
    return written
